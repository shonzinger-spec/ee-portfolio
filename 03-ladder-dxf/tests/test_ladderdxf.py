"""ladderdxf tests: the sample end to end, numbering rules, input checks, determinism, rendering."""

import csv
import os
import subprocess
import sys
from pathlib import Path

import ezdxf
import pytest
from openpyxl import Workbook, load_workbook

from ladderdxf import cli
from ladderdxf.drawing import TitleInfo, build_document
from ladderdxf.model import InputError, load_ladder
from ladderdxf.reports import bom_rows

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "sample" / "motor_starter.csv"
FIXED = ["--date", "2026-10-27", "--render", "none"]
HEADER = "rung,device,kind,description,around"


def run_sample(out: Path) -> int:
    return cli.main([str(SAMPLE), "--out", str(out), *FIXED])


def attribs(doc, tag: str) -> list[str]:
    return [a.dxf.text for ins in doc.modelspace().query("INSERT") for a in ins.attribs if a.dxf.tag == tag]


def write_csv(path: Path, rows: list[str]) -> Path:
    path.write_text("\n".join([HEADER, *rows]) + "\n", encoding="utf-8")
    return path


def sheet_rows(path: Path) -> list[tuple]:
    return list(load_workbook(path).active.iter_rows(min_row=2, values_only=True))


def dump(path: Path) -> list:
    """Type, layer, position and text of every model-space entity, for layout comparisons."""
    rows = []
    for e in ezdxf.readfile(path).modelspace():
        row = [e.dxftype(), e.dxf.layer]
        if e.dxftype() == "LINE":
            row += [round(v, 6) for v in (e.dxf.start.x, e.dxf.start.y, e.dxf.end.x, e.dxf.end.y)]
        elif e.dxftype() == "INSERT":
            row += [e.dxf.name, round(e.dxf.insert.x, 6), round(e.dxf.insert.y, 6)]
            row += [(a.dxf.tag, a.dxf.text, round(a.dxf.insert.x, 6), round(a.dxf.insert.y, 6)) for a in e.attribs]
        elif e.dxftype() == "TEXT":
            row += [e.dxf.text, round(e.dxf.insert.x, 6), round(e.dxf.insert.y, 6)]
        rows.append(row)
    return rows


@pytest.fixture(scope="module")
def sample_out(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("sample")
    assert run_sample(out) == 0
    return out


def test_dxf_reopens_with_unique_device_tags(sample_out):
    doc = ezdxf.readfile(sample_out / "motor_starter.dxf")
    tags = attribs(doc, "TAG1")
    assert sorted(tags) == ["LT1", "LT2", "M1", "OL1", "PB1", "PB2"]
    assert len(set(tags)) == len(tags)
    assert attribs(doc, "TAG2") == ["M1", "M1", "M1"]  # contacts carry their coil's tag
    assert not doc.audit().has_errors


def test_wire_numbers_are_unique_and_match_the_wire_list(sample_out):
    numbers = attribs(ezdxf.readfile(sample_out / "motor_starter.dxf"), "WIRENO")
    assert sorted(numbers, key=int) == [str(n) for n in range(1, 8)]
    rows = sheet_rows(sample_out / "motor_starter_wirelist.xlsx")
    assert {row[0] for row in rows} == set(numbers)
    assert [row[1:] for row in rows if row[0] == "4"] == [("PB2", "4", "M1", "A1"), ("M1", "A1", "M1", "14")]


def test_bom_has_one_row_per_input_device(sample_out):
    with SAMPLE.open(newline="", encoding="utf-8") as fh:
        devices = {row["device"] for row in csv.DictReader(fh)}
    rows = sheet_rows(sample_out / "motor_starter_bom.xlsx")
    assert len(rows) == len(devices) == 6
    assert [row[1] for row in rows] == ["LT1", "LT2", "M1", "OL1", "PB1", "PB2"]
    assert rows[2][2] == "Motor starter, coil + 2 NO / 1 NC aux"


def test_bom_merges_devices_with_the_same_part(tmp_path):
    ladder = load_ladder(write_csv(tmp_path / "lights.csv", [
        "1,START,PB_NO,START", "1,CR_A,RELAY_COIL,RELAY A",
        "2,CR_A,CONTACT_NO", "2,LAMP_A,PILOT_LIGHT,PUMP ON",
        "3,LAMP_B,PILOT_LIGHT,FAN ON",
        "4,LAMP_C,PILOT_LIGHT,HEAT ON",
    ]))
    for d, catalog in zip(ladder.devices, ["", "", "PL-22", "PL-22", ""]):
        d.parent.catalog, d.parent.manufacturer = catalog, "GENERIC" if catalog else ""
    rows = {row[1]: row for row in bom_rows(ladder)}
    assert set(rows) == {"CR1", "LT1, LT2", "LT3", "PB1"}  # LT3 has no catalog number, so it is not merged
    assert rows["LT1, LT2"][3:] == ("PUMP ON / FAN ON", "GENERIC", "PL-22", 2, "101,102")
    assert [row[0] for row in bom_rows(ladder)] == [1, 2, 3, 4]


def test_reports_are_named_after_the_input(tmp_path):
    other = write_csv(tmp_path / "pump.csv", ["1,START,PB_NO", "1,LAMP,PILOT_LIGHT"])
    assert cli.main([str(SAMPLE), "--out", str(tmp_path), *FIXED]) == 0
    assert cli.main([str(other), "--out", str(tmp_path), *FIXED]) == 0
    assert len(sheet_rows(tmp_path / "motor_starter_bom.xlsx")) == 6
    assert len(sheet_rows(tmp_path / "pump_bom.xlsx")) == 2
    assert (tmp_path / "pump_wirelist.xlsx").exists()


def test_cross_references(sample_out):
    doc = ezdxf.readfile(sample_out / "motor_starter.dxf")
    assert attribs(doc, "XREFNO") == ["NO 101,102"]
    assert attribs(doc, "XREFNC") == ["NC 103"]
    assert attribs(doc, "XREF") == ["100", "100", "100"]


def test_lines_and_wire_numbers_follow_rung_order():
    ladder = load_ladder(SAMPLE)
    assert [r.line for r in ladder.rungs] == [100, 102, 103]  # the seal-in branch takes line 101
    assert [r.nodes for r in ladder.rungs] == [["1", "3", "4", "5", "2"], ["1", "6", "2"], ["1", "7", "2"]]


def test_tags_count_per_family_in_drawing_order(tmp_path):
    ladder = load_ladder(write_csv(tmp_path / "relays.csv", [
        "1,START_A,PB_NO", "1,CR_A,RELAY_COIL",
        "2,START_B,PB_NO", "2,CR_B,RELAY_COIL",
        "3,CR_B,CONTACT_NO", "3,CR_A,CONTACT_NC", "3,LAMP,PILOT_LIGHT",
    ]))
    assert {d.key: d.tag for d in ladder.devices} == {"START_A": "PB1", "CR_A": "CR1", "START_B": "PB2", "CR_B": "CR2", "LAMP": "LT1"}
    first, second = ladder.rungs[2].main[:2]
    assert (first.tag, first.pins, first.xref) == ("CR2", ("13", "14"), "101")
    assert (second.tag, second.pins, second.xref) == ("CR1", ("11", "12"), "100")


def test_xlsx_input_gives_the_same_ladder(tmp_path):
    book = Workbook()
    with SAMPLE.open(newline="", encoding="utf-8") as fh:
        for i, row in enumerate(csv.reader(fh)):
            book.active.append([float(row[0]) if i else row[0], *row[1:]])  # Excel stores rung numbers as floats
    book.save(tmp_path / "motor_starter.xlsx")
    a, b = load_ladder(SAMPLE), load_ladder(tmp_path / "motor_starter.xlsx")
    assert [(d.key, d.tag) for d in a.devices] == [(d.key, d.tag) for d in b.devices]
    assert [(w.number, w.ends) for w in a.wires] == [(w.number, w.ends) for w in b.wires]
    assert bom_rows(a) == bom_rows(b)


def test_terminals_of_every_symbol_insert(sample_out):
    doc = ezdxf.readfile(sample_out / "motor_starter.dxf")
    terminals = list(zip(attribs(doc, "TERM01"), attribs(doc, "TERM02")))
    assert terminals == [
        ("1", "2"),    # PB1 STOP, NC pushbutton
        ("3", "4"),    # PB2 START, NO pushbutton
        ("A1", "A2"),  # M1 starter coil
        ("95", "96"),  # OL1 overload contact
        ("13", "14"),  # M1 seal-in NO contact around START, line 101
        ("23", "24"),  # M1 NO contact, line 102
        ("X1", "X2"),  # LT1, line 102 (inserts are drawn per rung: contact then light)
        ("31", "32"),  # M1 NC contact, line 103
        ("X1", "X2"),  # LT2, line 103
    ]


@pytest.mark.parametrize("rows, message", [
    (["1,STOP,PB_NC"], "no load"),
    (["1,LT_A,PILOT_LIGHT", "1,LT_B,PILOT_LIGHT"], "one load per rung"),
    (["1,X,CONTACT_NO", "1,LAMP,PILOT_LIGHT"], "has no coil"),
    (["1,X,FUSE", "1,LAMP,PILOT_LIGHT"], "unknown kind"),
    (["1,S,PB_NO", "1,LAMP,PILOT_LIGHT", "1,T,PB_NO,,NOPE"], "around='NOPE'"),
    (["1,LAMP,PILOT_LIGHT", "1,S,PB_NO"], "only overload contacts may follow the load"),
    (["1,OLX,OL_NC", "1,LAMP,PILOT_LIGHT"], "right of the load"),
    (["x,S,PB_NO", "x,LAMP,PILOT_LIGHT"], "whole number"),
    (["1,S,PB_NO", "1,LAMP,PILOT_LIGHT", "2,S,PB_NO", "2,LAMP2,PILOT_LIGHT"], "draws each pushbutton"),
    pytest.param(
        ["1,START,PB_NO", "1,CRX,STARTER_COIL", "1,CRX,CONTACT_NO,,CRX"],
        "go around a contact or pushbutton, not STARTER_COIL CRX",
        id="branch-around-a-coil",
    ),
    pytest.param(
        ["1,S,PB_NO", "1,CRX,STARTER_COIL", "1,LAMP,PILOT_LIGHT,,S"],
        "can be a branch, not PILOT_LIGHT LAMP",
        id="load-as-branch",
    ),
    pytest.param(
        ["1,S,PB_NO", "1,CRX,STARTER_COIL", "1,OLX,OL_NC,,S"],
        "can be a branch, not OL_NC OLX",
        id="overload-as-branch",
    ),
    pytest.param(
        ["1,S,PB_NO", "1,CRX,STARTER_COIL", "1,OLX,OL_NC", "1,CRX,CONTACT_NO,,OLX"],
        "go around a contact or pushbutton, not OL_NC OLX",
        id="branch-around-an-overload",
    ),
    pytest.param(["²,S,PB_NO", "²,LAMP,PILOT_LIGHT"], "whole number", id="superscript-rung-number"),
    pytest.param(
        ["1,START,PB_NO", "1,CRX,STARTER_COIL"]
        + [f"{n},CRX,CONTACT_NO" for n in range(2, 12)]
        + [f"{n},LT{n},PILOT_LIGHT" for n in range(2, 12)],
        "at most 9",
        id="ten-contacts-of-one-coil",
    ),
])
def test_bad_device_lists_are_rejected(tmp_path, rows, message):
    with pytest.raises(InputError, match=message):
        load_ladder(write_csv(tmp_path / "bad.csv", rows))


def test_missing_column_is_reported(tmp_path):
    path = tmp_path / "nokind.csv"
    path.write_text("rung,device\n1,STOP\n", encoding="utf-8")
    with pytest.raises(InputError, match="missing column"):
        load_ladder(path)


def test_non_utf8_csv_is_reported(tmp_path):
    path = tmp_path / "heater.csv"
    path.write_text(HEADER + "\n1,HEAT_LT,PILOT_LIGHT,SETPOINT 80°C\n", encoding="cp1252")
    with pytest.raises(InputError, match="UTF-8"):
        load_ladder(path)


def test_corrupt_xlsx_is_reported(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"PK\x03\x04 this is not a workbook")
    with pytest.raises(InputError, match="Excel workbook"):
        load_ladder(path)


def test_more_lines_than_one_sheet_is_rejected(tmp_path):
    ladder = load_ladder(write_csv(tmp_path / "long.csv", [f"{n},LAMP{n},PILOT_LIGHT" for n in range(1, 12)]))
    with pytest.raises(InputError, match="single sheet"):
        build_document(ladder, TitleInfo())


def test_too_many_symbols_on_one_line_is_rejected(tmp_path):
    rows = [f"1,PB{n},PB_NO" for n in range(8)] + ["1,LAMP,PILOT_LIGHT"]
    with pytest.raises(InputError, match="do not fit"):
        build_document(load_ladder(write_csv(tmp_path / "wide.csv", rows)), TitleInfo())


def test_cli_reports_input_errors_without_a_traceback(tmp_path, capsys):
    assert cli.main([str(write_csv(tmp_path / "bad.csv", ["1,STOP,PB_NC"])), "--out", str(tmp_path), *FIXED]) == 2
    assert "no load" in capsys.readouterr().err


def test_layout_is_the_same_in_a_fresh_process(tmp_path):
    assert run_sample(tmp_path / "a") == 0
    env = dict(os.environ, PYTHONHASHSEED="12345")
    proc = subprocess.run(
        [sys.executable, "-m", "ladderdxf", str(SAMPLE), "--out", str(tmp_path / "b"), *FIXED],
        cwd=ROOT, env=env, capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert dump(tmp_path / "a" / "motor_starter.dxf") == dump(tmp_path / "b" / "motor_starter.dxf")


def test_render_writes_a_png(tmp_path):
    pytest.importorskip("matplotlib")
    assert cli.main([str(SAMPLE), "--out", str(tmp_path), "--date", "2026-10-27", "--render", "png"]) == 0
    png = tmp_path / "motor_starter.png"
    assert png.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
    assert png.stat().st_size > 20_000
