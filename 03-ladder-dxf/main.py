#!/usr/bin/env python3
"""Project 03 standalone run: draw the bundled motor-starter sample and check what was written.

    pip install -r requirements.txt
    python main.py

Writes out/motor_starter.dxf, out/motor_starter.png, out/bom.xlsx and out/wirelist.xlsx, then
re-opens them and prints the checks quoted in the README's Verification section.
"""

from __future__ import annotations

import csv
import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import ezdxf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from ladderdxf.cli import main as ladderdxf  # noqa: E402

SAMPLE = Path("sample") / "motor_starter.csv"  # relative to HERE; main() changes into it
OUT = Path("out")


def checks() -> list[tuple[bool, str]]:
    doc = ezdxf.readfile(OUT / "motor_starter.dxf")
    found: dict[str, list[str]] = {"TAG1": [], "TAG2": [], "WIRENO": []}
    for insert in doc.modelspace().query("INSERT"):
        for attrib in insert.attribs:
            if attrib.dxf.tag in found:
                found[attrib.dxf.tag].append(attrib.dxf.text)
    tags, contacts, wires = found["TAG1"], found["TAG2"], found["WIRENO"]
    with SAMPLE.open(newline="", encoding="utf-8-sig") as fh:
        devices = {row["device"].strip() for row in csv.DictReader(fh)}
    bom = list(load_workbook(OUT / "bom.xlsx").active.iter_rows(min_row=2, values_only=True))
    wire_list = {str(row[0]) for row in load_workbook(OUT / "wirelist.xlsx").active.iter_rows(min_row=2, values_only=True)}
    audit = doc.audit()
    return [
        (len(set(tags)) == len(tags) == len(devices), f"device tags unique: {len(set(tags))} of {len(tags)} -> {' '.join(sorted(tags))}"),
        (bool(contacts) and set(contacts) <= set(tags), f"all {len(contacts)} contact tags belong to a drawn device -> {' '.join(sorted(set(contacts)))}"),
        (len(set(wires)) == len(wires), f"wire numbers unique: {len(set(wires))} of {len(wires)} -> {' '.join(sorted(wires, key=int))}"),
        (len(bom) == len(devices), f"BOM rows {len(bom)} = devices in the input list {len(devices)}"),
        (wire_list == set(wires), f"wire list covers the same {len(wire_list)} wire numbers as the drawing"),
        (not audit.has_errors, f"ezdxf audit of the DXF: {len(audit.errors)} errors"),
    ]


def main() -> int:
    os.chdir(HERE)  # relative paths keep the printed output short
    code = ladderdxf([
        str(SAMPLE), "--out", str(OUT), "--render", "png",
        "--project", "EE PORTFOLIO 03 - LADDER-DXF", "--title", "MOTOR 1 STARTER", "--drawn-by", "S. ZINGER",
    ])
    if code:
        return code
    results = checks()
    print("checks:")
    for ok, text in results:
        print(f"  [{'PASS' if ok else 'FAIL'}] {text}")
    passed = all(ok for ok, _ in results)
    print("ALL CHECKS PASSED" if passed else "SOME CHECKS FAILED")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
