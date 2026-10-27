"""Device list -> ladder model: input checks, line references, device tags, wire numbers, cross-references.

Geometry lives in drawing.py; this module only decides what goes on the sheet.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

FIRST_LINE = 100  # first line reference, as on sheet 1 of an AutoCAD Electrical ladder
RAILS = (("L1", "1"), ("L2", "2"))  # (rail name, wire number): left = 120 VAC hot, right = neutral
REQUIRED_COLUMNS = ("rung", "device", "kind")
COLOR_NAMES = {"R": "red", "G": "green", "A": "amber", "Y": "yellow", "B": "blue", "W": "white", "C": "clear"}


class InputError(ValueError):
    """The device list cannot be drawn; the message names every problem found."""


@dataclass(frozen=True)
class Kind:
    family: str  # tag prefix; empty for contacts, which carry the tag of their coil's device
    role: str  # "contact", "load" or "overload"
    block: str  # DXF block that draws it
    pins: tuple[str, str] = ("", "")  # terminal numbers (contacts get theirs from their coil)
    label: str = ""  # BOM wording when this element is the device's main part


KINDS = {
    "PB_NO": Kind("PB", "contact", "NEMA_PB_NO", ("3", "4"), "Pushbutton, 1 NO"),
    "PB_NC": Kind("PB", "contact", "NEMA_PB_NC", ("1", "2"), "Pushbutton, 1 NC"),
    "CONTACT_NO": Kind("", "contact", "NEMA_CONTACT_NO"),
    "CONTACT_NC": Kind("", "contact", "NEMA_CONTACT_NC"),
    "OL_NC": Kind("OL", "overload", "NEMA_OL_NC", ("95", "96"), "Overload relay"),
    "STARTER_COIL": Kind("M", "load", "NEMA_COIL", ("A1", "A2"), "Motor starter"),
    "RELAY_COIL": Kind("CR", "load", "NEMA_COIL", ("A1", "A2"), "Control relay"),
    "PILOT_LIGHT": Kind("LT", "load", "NEMA_PILOT_LIGHT", ("X1", "X2"), "Pilot light"),
}
COILS = frozenset({"STARTER_COIL", "RELAY_COIL"})
CONTACTS = frozenset({"CONTACT_NO", "CONTACT_NC"})


@dataclass(eq=False)
class Element:
    row: int  # row number in the input file (the header is row 1)
    rung: int
    device: str
    kind: str
    description: str = ""
    around: str = ""
    color: str = ""
    manufacturer: str = ""
    catalog: str = ""
    line: int = 0  # line reference on the sheet
    pos: int = 0  # slot on its line; a branch element takes the slot of the element it bridges
    tag: str = ""
    pins: tuple[str, str] = ("", "")
    xref: str = ""  # contacts only: line reference of their coil

    @property
    def spec(self) -> Kind:
        return KINDS[self.kind]

    @property
    def color_letter(self) -> str:
        return self.color.strip()[:1].upper()


@dataclass(eq=False)
class Rung:
    number: int
    main: list[Element]  # left to right between the rails
    branches: list[Element]  # each on its own line below, bridging one main element
    targets: list[int] = field(default_factory=list)  # index in `main` bridged by each branch
    load_index: int = 0
    line: int = 0
    nodes: list[str] = field(default_factory=list)  # wire number of each node, rail to rail


@dataclass(eq=False)
class Device:
    key: str
    parent: Element  # the coil, or the only element of a pushbutton / overload / light
    elements: list[Element]
    tag: str = ""
    number: int = 0
    xref_no: list[int] = field(default_factory=list)  # lines of this coil's NO contacts
    xref_nc: list[int] = field(default_factory=list)  # lines of this coil's NC contacts

    @property
    def family(self) -> str:
        return self.parent.spec.family

    @property
    def label(self) -> str:
        base = self.parent.spec.label
        if self.parent.kind in COILS:
            return f"{base}, coil + {len(self.xref_no)} NO / {len(self.xref_nc)} NC aux"
        if self.parent.kind == "PILOT_LIGHT" and self.parent.color:
            return f"{base}, {COLOR_NAMES.get(self.parent.color_letter, self.parent.color.lower())}"
        return base


@dataclass(eq=False)
class Wire:
    number: str
    ends: list[tuple[str, str]]  # (tag, pin) in drawing order; a rail starts at its source name


@dataclass(eq=False)
class Ladder:
    rungs: list[Rung]
    devices: list[Device]  # in tag order, which is drawing order
    wires: list[Wire]  # rails first, then the nodes in rung order

    @property
    def last_line(self) -> int:
        return max(e.line for r in self.rungs for e in (*r.main, *r.branches))

    def device(self, key: str) -> Device:
        return next(d for d in self.devices if d.key == key)


def load_ladder(path) -> Ladder:
    return build_ladder(read_rows(path), Path(path).name)


def read_rows(path) -> list[tuple[int, dict[str, str]]]:
    """(row number, {column: text}) for each non-blank row of a .csv or the first sheet of an .xlsx."""
    path = Path(path)
    if not path.is_file():
        raise InputError(f"{path}: file not found")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            table = [list(row) for row in csv.reader(fh)]
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            table = [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
        finally:
            book.close()
    else:
        raise InputError(f"{path.name}: expected a .csv or .xlsx device list")
    if not table:
        raise InputError(f"{path.name}: file is empty")
    header = [_text(h).lower().replace(" ", "_") for h in table[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise InputError(f"{path.name}: missing column(s): {', '.join(missing)}")
    rows = []
    for number, values in enumerate(table[1:], start=2):
        cells = [_text(v) for v in values]
        if any(cells):
            rows.append((number, {name: cells[i] if i < len(cells) else "" for i, name in enumerate(header) if name}))
    return rows


def build_ladder(rows, source: str = "device list") -> Ladder:
    errors: list[str] = []
    elements = [el for el in (_element(n, row, errors) for n, row in rows) if el]
    if not rows:
        errors.append("no device rows")
    _fail(errors, source)

    by_rung: dict[int, list[Element]] = {}
    for el in elements:
        by_rung.setdefault(el.rung, []).append(el)
    rungs = [_rung(number, by_rung[number], errors) for number in sorted(by_rung)]
    devices = _devices(elements, errors)
    _fail(errors, source)

    _number_lines(rungs)
    devices = _tag_devices(devices)
    _link_contacts(devices)
    return Ladder(rungs, devices, _number_wires(rungs))


def _text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 1 as 1.0
    return "" if value is None else str(value).strip()


def _fail(errors: list[str], source: str) -> None:
    if errors:
        raise InputError(f"{source} has {len(errors)} problem(s):\n" + "\n".join(f"  - {e}" for e in errors))


def _element(number: int, row: dict[str, str], errors: list[str]) -> Element | None:
    rung, device, kind = row.get("rung", ""), row.get("device", ""), row.get("kind", "").upper()
    if not rung.isdigit() or int(rung) < 1:
        errors.append(f"row {number}: rung must be a whole number from 1 up, got {rung!r}")
    elif not device:
        errors.append(f"row {number}: device is empty")
    elif kind not in KINDS:
        errors.append(f"row {number}: unknown kind {row.get('kind', '')!r}; use one of {', '.join(KINDS)}")
    else:
        extra = (row.get(c, "") for c in ("description", "around", "color", "manufacturer", "catalog"))
        return Element(number, int(rung), device, kind, *extra)
    return None


def _rung(number: int, elements: list[Element], errors: list[str]) -> Rung:
    rung = Rung(number, [e for e in elements if not e.around], [e for e in elements if e.around])
    where = f"rung {number}"
    loads = [i for i, e in enumerate(rung.main) if e.spec.role == "load"]
    if not rung.main:
        errors.append(f"{where}: every row has 'around', so there is no main path")
    elif not loads:
        errors.append(f"{where}: no load (coil or pilot light), so it would short L1 to L2")
    elif len(loads) > 1:
        names = ", ".join(rung.main[i].device for i in loads)
        errors.append(f"{where}: loads {names} are in series; NEMA allows one load per rung")
    else:
        rung.load_index = loads[0]
        for e in rung.main[: rung.load_index]:
            if e.spec.role == "overload":
                errors.append(f"row {e.row}: overload contact {e.device} goes right of the load (NEMA convention)")
        for e in rung.main[rung.load_index + 1 :]:
            if e.spec.role != "overload":
                errors.append(f"row {e.row}: only overload contacts may follow the load, not {e.kind} {e.device}")
    for b in rung.branches:
        hits = [i for i, e in enumerate(rung.main) if e.device == b.around]
        if len(hits) == 1:
            rung.targets.append(hits[0])
        else:
            errors.append(f"row {b.row}: around={b.around!r} must name one main-path device on {where}, found {len(hits)}")
    return rung


def _devices(elements: list[Element], errors: list[str]) -> list[Device]:
    by_key: dict[str, list[Element]] = {}
    for el in elements:
        by_key.setdefault(el.device, []).append(el)
    devices = []
    for key, els in by_key.items():
        coils = [e for e in els if e.kind in COILS]
        rows = ", ".join(str(e.row) for e in els)
        if coils:
            if len(coils) > 1:
                errors.append(f"device {key}: {len(coils)} coils (rows {rows}); a device has one coil")
            for e in els:
                if e.kind not in COILS | CONTACTS:
                    errors.append(f"row {e.row}: {key} has a coil, so its other rows must be CONTACT_NO or CONTACT_NC, not {e.kind}")
        else:
            for e in els:
                if e.kind in CONTACTS:
                    errors.append(f"row {e.row}: {e.kind} of {key}, but {key} has no coil row (STARTER_COIL or RELAY_COIL)")
            if len(els) > 1 and not any(e.kind in CONTACTS for e in els):
                errors.append(f"device {key}: listed on rows {rows}; v0 draws each pushbutton, overload and light once")
        devices.append(Device(key, coils[0] if coils else els[0], els))
    return devices


def _number_lines(rungs: list[Rung]) -> None:
    """Main path of each rung on one line, then one line per branch, numbered from FIRST_LINE."""
    line = FIRST_LINE
    for rung in rungs:
        rung.line = line
        for pos, e in enumerate(rung.main):
            e.line, e.pos = line, pos
        for b, target in zip(rung.branches, rung.targets):
            line += 1
            b.line, b.pos = line, target
        line += 1


def _tag_devices(devices: list[Device]) -> list[Device]:
    """Sequential tags per family (PB1, PB2, M1, ...) in drawing order of each device's main part."""
    ordered = sorted(devices, key=lambda d: (d.parent.line, d.parent.pos))
    counts: dict[str, int] = {}
    for d in ordered:
        counts[d.family] = d.number = counts.get(d.family, 0) + 1
        d.tag = f"{d.family}{d.number}"
        for e in d.elements:
            e.tag = d.tag
    return ordered


def _link_contacts(devices: list[Device]) -> None:
    """Terminal numbers for every element; contacts also get their coil's line, coils their contacts' lines."""
    for d in devices:
        for e in d.elements:
            e.pins = e.spec.pins
        if d.parent.kind in COILS:
            contacts = sorted((e for e in d.elements if e.kind in CONTACTS), key=lambda e: (e.line, e.pos))
            for n, c in enumerate(contacts, start=1):  # IEC/EN 50011 style: 13-14 first NO, 21-22 second (NC)...
                c.pins = (f"{n}3", f"{n}4") if c.kind == "CONTACT_NO" else (f"{n}1", f"{n}2")
                c.xref = str(d.parent.line)
            d.xref_no = [c.line for c in contacts if c.kind == "CONTACT_NO"]
            d.xref_nc = [c.line for c in contacts if c.kind == "CONTACT_NC"]


def _number_wires(rungs: list[Rung]) -> list[Wire]:
    """Rails are wires 1 (L1) and 2 (L2); every node between two main-path symbols gets 3, 4, 5, ...

    A branch bridges two existing nodes, so it adds terminals to them but no new wire.
    """
    (left_name, left_no), (right_name, right_no) = RAILS
    left, right = Wire(left_no, [(left_name, "")]), Wire(right_no, [(right_name, "")])
    nodes: list[Wire] = []
    for rung in rungs:
        inner = [Wire(str(3 + len(nodes) + k), []) for k in range(len(rung.main) - 1)]
        nodes += inner
        wires = [left, *inner, right]  # wires[k] is the node left of rung.main[k]
        rung.nodes = [w.number for w in wires]
        for k, e in enumerate(rung.main):
            wires[k].ends.append((e.tag, e.pins[0]))
            wires[k + 1].ends.append((e.tag, e.pins[1]))
        for b, k in zip(rung.branches, rung.targets):
            wires[k].ends.append((b.tag, b.pins[0]))
            wires[k + 1].ends.append((b.tag, b.pins[1]))
    return [left, right, *nodes]
