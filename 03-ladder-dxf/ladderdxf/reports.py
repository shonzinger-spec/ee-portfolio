"""BOM and from-to wire list as XLSX (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .model import Ladder

BOM_HEADER = ("Item", "Tag", "Device", "Description", "Manufacturer", "Catalog", "Qty", "Lines")
WIRE_HEADER = ("Wire", "From", "From pin", "To", "To pin")


def bom_rows(ladder: Ladder) -> list[tuple]:
    """One row per device (tag), sorted by tag family, then number."""
    devices = sorted(ladder.devices, key=lambda d: (d.family, d.number))
    return [
        (item, d.tag, d.label, d.parent.description, d.parent.manufacturer, d.parent.catalog, 1,
         ",".join(str(n) for n in sorted({e.line for e in d.elements})))
        for item, d in enumerate(devices, start=1)
    ]


def wire_rows(ladder: Ladder) -> list[tuple]:
    """From-to rows: each wire daisy-chains its terminals in drawing order; a rail starts at L1 / L2."""
    return [(w.number, *a, *b) for w in ladder.wires for a, b in zip(w.ends, w.ends[1:])]


def write_bom(ladder: Ladder, path) -> Path:
    return _write(path, "BOM", BOM_HEADER, bom_rows(ladder))


def write_wire_list(ladder: Ladder, path) -> Path:
    return _write(path, "Wire list", WIRE_HEADER, wire_rows(ladder))


def _write(path, title: str, header: tuple, rows: list[tuple]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for i, column in enumerate(zip(header, *rows), start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(len(str(v)) for v in column) + 2
    book.save(path)
    return Path(path)
